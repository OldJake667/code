"""Build a concise presentation focused only on validation accuracy results."""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)

# Best saved-checkpoint results from the completed notebook/local runs.
RESULTS = [
    ("CNN", "From scratch", 42),
    ("ViT-Tiny", "Frozen", 42),
    ("ViT-Tiny", "Fine-tuned", 43),
    ("DINOv2", "Frozen", 43),
    ("DINOv2", "Fine-tuned", 43),
    ("VMamba", "Frozen", 42),
    ("VMamba", "Fine-tuned", 44),
]
VALIDATION_SIZE = 45

BG = "F5F7FA"
NAVY = "152536"
TEXT = "344657"
MUTED = "667788"
BLUE = "28689D"
TEAL = "188C87"
ORANGE = "E1843B"
WHITE = "FFFFFF"


def accuracy(correct):
    return 100 * correct / VALIDATION_SIZE


def make_overall_chart():
    labels = [
        "CNN",
        "ViT-Tiny\nfrozen",
        "ViT-Tiny\nfine-tuned",
        "DINOv2\nfrozen",
        "DINOv2\nfine-tuned",
        "VMamba\nfrozen",
        "VMamba\nfine-tuned",
    ]
    values = [accuracy(correct) for _, _, correct in RESULTS]
    colors = [BLUE, BLUE, TEAL, BLUE, TEAL, BLUE, TEAL]

    figure, axis = plt.subplots(figsize=(12.4, 5.3))
    bars = axis.bar(range(len(values)), values, color=[f"#{value}" for value in colors], width=.63)
    axis.set_ylim(88, 100)
    axis.set_ylabel("Validation accuracy (%)", fontsize=12)
    axis.set_xticks(range(len(labels)), labels, fontsize=10)
    axis.set_yticks([88, 90, 92, 94, 96, 98, 100])
    axis.grid(axis="y", alpha=.2)
    axis.set_axisbelow(True)
    axis.spines[["top", "right", "left"]].set_visible(False)
    axis.tick_params(axis="y", length=0)
    for bar, value, (_, _, correct) in zip(bars, values, RESULTS):
        axis.text(
            bar.get_x() + bar.get_width() / 2,
            value + .3,
            f"{value:.2f}%\n({correct}/45)",
            ha="center",
            va="bottom",
            fontsize=10,
            fontweight="bold",
            color=f"#{NAVY}",
        )
    figure.tight_layout()
    path = OUT / "validation_accuracy_all_models.png"
    figure.savefig(path, dpi=220, facecolor="white", bbox_inches="tight")
    plt.close(figure)
    return path


def make_pair_chart():
    models = ["ViT-Tiny", "DINOv2", "VMamba"]
    frozen = [accuracy(42), accuracy(43), accuracy(42)]
    tuned = [accuracy(43), accuracy(43), accuracy(44)]
    x = range(len(models))

    figure, axis = plt.subplots(figsize=(11.8, 5.2))
    width = .3
    frozen_bars = axis.bar([i - width / 2 for i in x], frozen, width, label="Frozen", color=f"#{BLUE}")
    tuned_bars = axis.bar([i + width / 2 for i in x], tuned, width, label="Fine-tuned", color=f"#{TEAL}")
    axis.set_ylim(88, 100)
    axis.set_ylabel("Validation accuracy (%)", fontsize=12)
    axis.set_xticks(list(x), models, fontsize=12)
    axis.set_yticks([88, 90, 92, 94, 96, 98, 100])
    axis.grid(axis="y", alpha=.2)
    axis.set_axisbelow(True)
    axis.spines[["top", "right", "left"]].set_visible(False)
    axis.tick_params(axis="y", length=0)
    axis.legend(frameon=False, loc="upper left", ncol=2)
    for bars in (frozen_bars, tuned_bars):
        for bar in bars:
            axis.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + .25,
                f"{bar.get_height():.2f}%",
                ha="center",
                va="bottom",
                fontsize=10,
                fontweight="bold",
                color=f"#{NAVY}",
            )
    figure.tight_layout()
    path = OUT / "validation_accuracy_frozen_vs_finetuned.png"
    figure.savefig(path, dpi=220, facecolor="white", bbox_inches="tight")
    plt.close(figure)
    return path


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def new_slide(prs, title, eyebrow="VALIDATION RESULTS"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(BG)
    add_text(slide, eyebrow, .68, .22, 12, .3, 10, TEAL, True)
    add_text(slide, title, .68, .55, 12, .65, 28, NAVY, True)
    return slide


def add_footer(slide, value):
    add_text(slide, value, .7, 7.08, 11.95, .25, 10, MUTED, False, PP_ALIGN.CENTER)


overall_chart = make_overall_chart()
pair_chart = make_pair_chart()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = rgb(NAVY)
add_text(slide, "PAPER REPLICATION", .78, .62, 11.8, .35, 11, "6DD1CA", True)
add_text(slide, "Validation accuracy results", .78, 1.45, 11.8, .9, 34, WHITE, True)
add_text(slide, "CNN, ViT-Tiny, DINOv2 and VMamba", .8, 2.55, 11.7, .55, 22, "B9D4E5", False)
add_text(slide, "Best saved checkpoints on the shared 45-image validation split", .8, 5.75, 11.7, .5, 17, WHITE, False)
add_text(slide, "31 August 2026", .8, 6.38, 11.7, .35, 12, "A6B8C8", False)

# Overall comparison
slide = new_slide(prs, "Best validation accuracy by model")
slide.shapes.add_picture(str(overall_chart), Inches(.62), Inches(1.27), width=Inches(12.1))
add_footer(slide, "All values use the best saved validation checkpoint. One image equals 2.22 percentage points.")

# Frozen versus fine-tuned
slide = new_slide(prs, "Effect of fine-tuning the pretrained models")
slide.shapes.add_picture(str(pair_chart), Inches(.75), Inches(1.3), width=Inches(11.85))
add_footer(slide, "Fine-tuning changed ViT-Tiny by +2.22 pp, DINOv2 by 0.00 pp, and VMamba by +4.44 pp.")

# Takeaway
slide = new_slide(prs, "What the comparison shows", "INTERPRETATION")
cards = [
    ("44/45", "97.78%", "VMamba fine-tuned", TEAL),
    ("43/45", "95.56%", "ViT-Tiny FT and both DINOv2 runs", BLUE),
    ("42/45", "93.33%", "CNN, ViT-Tiny frozen and VMamba frozen", ORANGE),
]
for index, (correct, percent, label, accent) in enumerate(cards):
    x = .75 + index * 4.08
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(1.55), Inches(3.7), Inches(3.75),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = rgb(WHITE)
    box.line.color.rgb = rgb("D8E0E7")
    strip = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x), Inches(1.55), Inches(.08), Inches(3.75),
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = rgb(accent)
    strip.line.fill.background()
    add_text(slide, correct, x + .28, 1.92, 3.1, .65, 31, NAVY, True)
    add_text(slide, percent, x + .28, 2.62, 3.1, .45, 18, accent, True)
    add_text(slide, label, x + .28, 3.43, 3.05, 1.05, 15, TEXT, False)
add_text(
    slide,
    "These are validation results, not independent test results. The small split supports a controlled comparison, but not a claim of general architectural superiority.",
    .9, 5.82, 11.5, .78, 16, MUTED, False, PP_ALIGN.CENTER,
)

output = OUT / "validation_accuracy_results_20260831.pptx"
prs.save(output)
print(output)
