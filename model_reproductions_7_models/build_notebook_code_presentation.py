"""Create an editable presentation from the local notebooks, results and real model code."""

import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)


# The reported result follows the saved best validation checkpoint, because that
# is the checkpoint loaded back into the model for the final reported accuracy.
SOURCES = [
    ("CNN", None, 42 / 45, 42 / 45, "Best checkpoint reported"),
    ("ViT-Tiny", ROOT / "02_vit_tiny_frozen/notebook_results_run2", None, None, "Notebook rerun"),
    ("ViT-Tiny FT", ROOT / "03_vit_tiny_fine_tuned/notebook_results_cpu_notebook", None, None, "CPU-stable notebook rerun"),
    ("DINOv2", ROOT / "04_dinov2_frozen/notebook_results", None, None, "Notebook rerun"),
    ("DINOv2 FT", ROOT / "05_dinov2_fine_tuned/notebook_results_cpu_notebook", None, None, "CPU-stable notebook rerun"),
    ("VMamba", ROOT / "06_vmamba_frozen/notebook_results", None, None, "Notebook rerun"),
    ("VMamba FT", ROOT / "07_vmamba_fine_tuned/notebook_results_fixed_cpu", None, None, "CPU-stable local run; notebook in progress"),
]


def load_rows():
    rows = []
    for label, folder, final, best, status in SOURCES:
        if folder is not None:
            summary = json.loads((folder / "result_summary.json").read_text())
            best = summary["best_accuracy"]["validation_accuracy"]
            final = best
        rows.append({"label": label, "final": final, "best": best, "status": status})
    return rows


ROWS = load_rows()


def make_bar_chart():
    labels = [row["label"] for row in ROWS]
    x = list(range(len(labels)))
    figure, axis = plt.subplots(figsize=(12, 5.0))
    reported = [100 * row["best"] for row in ROWS]
    bars = axis.bar(x, reported, 0.52, color="#27649b")
    axis.set_ylim(84, 101.5)
    axis.set_ylabel("Validation accuracy (%)")
    axis.set_xticks(x, labels)
    axis.grid(axis="y", alpha=.22)
    for bar in bars:
        axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + .28,
                  f"{bar.get_height():.2f}", ha="center", fontsize=9)
    figure.tight_layout()
    path = OUT / "notebook_code_accuracy_comparison.png"
    figure.savefig(path, dpi=220, facecolor="white")
    plt.close(figure)
    return path


BAR_CHART = make_bar_chart()


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

BG, WHITE, NAVY, TEXT, MUTED, BLUE, TEAL, ORANGE = (
    "F3F7FA", "FFFFFF", "0E1F32", "263746", "607182", "25639B", "1B9190", "E17E2D"
)


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, value, x, y, w, h, size=18, color=TEXT, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.04)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    run = paragraph.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return shape


def slide(title, eyebrow):
    page = prs.slides.add_slide(blank)
    background = page.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid(); background.fill.fore_color.rgb = rgb(BG); background.line.fill.background()
    page.shapes._spTree.remove(background._element); page.shapes._spTree.insert(2, background._element)
    text(page, eyebrow.upper(), .68, .22, 12, .25, 10, TEAL, True)
    text(page, title, .68, .55, 12, .58, 27, NAVY, True)
    return page


def card(page, x, y, w, h, title, body, accent=BLUE):
    box = page.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = rgb(WHITE); box.line.color.rgb = rgb("D9E2E8")
    line = page.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(.075), Inches(h))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(accent); line.line.fill.background()
    text(page, title, x+.24, y+.2, w-.45, .35, 16, NAVY, True)
    text(page, body, x+.24, y+.67, w-.45, h-.86, 13, TEXT)


def code_box(page, title, code, x, y, w, h, accent=BLUE):
    card(page, x, y, w, h, title, "", accent)
    text(page, code, x+.28, y+.7, w-.5, h-.9, 10.5, "253746", False, "Menlo")


def add_curve(page, folder, x=.65, y=1.37, w=6.1):
    image = ROOT / folder / "training_history.png"
    if image.is_file():
        page.shapes.add_picture(str(image), Inches(x), Inches(y), width=Inches(w))
    else:
        card(page, x, y, w, 4.7, "Curve unavailable", "This notebook run did not complete with valid metrics.", ORANGE)


def add_image(page, path, x=.65, y=1.37, w=6.1):
    image = ROOT / path
    if image.is_file():
        page.shapes.add_picture(str(image), Inches(x), Inches(y), width=Inches(w))
    else:
        card(page, x, y, w, 4.7, "Image missing", str(path), ORANGE)


# Slide 1
page = slide("Model results and code notes", "AI models")
text(page, "CNN · ViT-Tiny · DINOv2 · VMamba", .72, 1.75, 11.8, .5, 24, BLUE, True)
card(page, .75, 4.25, 3.7, 1.35, "Data", "Same 45 validation images.", BLUE)
card(page, 4.82, 4.25, 3.7, 1.35, "Metric", "Best saved checkpoint.", TEAL)
card(page, 8.89, 4.25, 3.7, 1.35, "Code", "Real local code snippets.", ORANGE)

# Slide 2
page = slide("Same-split accuracy results", "results")
page.shapes.add_picture(str(BAR_CHART), Inches(.6), Inches(1.28), width=Inches(12.1))
text(page, "Reported accuracy uses the best saved checkpoint. One image = 2.22 percentage points.",
     .72, 6.73, 11.8, .35, 12, MUTED, True, align=PP_ALIGN.CENTER)

# Slide 3
page = slide("CNN from scratch", "CNN")
add_image(page, "01_cnn_from_scratch/training_history.png")
code_box(page, "CNN model", "model = Sequential([\n    Conv2D(32, (3, 3), activation=\"relu\"),\n    MaxPooling2D((2, 2)),\n    Conv2D(64, (3, 3), activation=\"relu\"),\n    MaxPooling2D((2, 2)),\n    Conv2D(128, (3, 3), activation=\"relu\"),\n    MaxPooling2D((2, 2)),\n    Flatten(),\n    Dense(128, activation=\"relu\"),\n    Dropout(0.5),\n    Dense(2, activation=\"softmax\"),\n])", 7.0, 1.48, 5.65, 4.55, TEAL)
text(page, "Desktop-slide result used in the chart: 42/45 = 93.33%.", 7.08, 6.18, 5.4, .45, 15, NAVY, True)

# Slide 4
page = slide("ViT-Tiny frozen", "Vision Transformer")
add_curve(page, "02_vit_tiny_frozen/notebook_results_run2")
code_box(page, "Model code", "class ViTTiny(nn.Module):\n    def __init__(self, num_classes=2):\n        self.patch_embed = PatchEmbedding(16, 192)\n        self.cls_token = nn.Parameter(...)\n        self.pos_embed = nn.Parameter(...)\n        self.blocks = nn.ModuleList([\n            TransformerBlock(192, 3) for _ in range(12)\n        ])\n        self.norm = nn.LayerNorm(192)\n        self.head = nn.Linear(192, num_classes)", 7.0, 1.48, 5.65, 4.25, TEAL)
text(page, "Best checkpoint: 42/45 = 93.33%.", 7.08, 6.05, 5.4, .45, 16, NAVY, True)

# Slide 5
page = slide("ViT-Tiny fine-tuning", "Vision Transformer")
add_curve(page, "03_vit_tiny_fine_tuned/notebook_results_cpu_notebook")
code_box(page, "Freeze / unfreeze", "for parameter in model.parameters():\n    parameter.requires_grad = False\n\nfor parameter in model.blocks[-1].parameters():\n    parameter.requires_grad = True\n\nfor parameter in model.head.parameters():\n    parameter.requires_grad = True", 7.0, 1.55, 5.65, 3.75, ORANGE)
text(page, "Best checkpoint: 43/45 = 95.56%.",
     7.08, 5.65, 5.4, .55, 15, NAVY, True)

# Slide 6
page = slide("DINOv2 frozen", "DINOv2")
add_curve(page, "04_dinov2_frozen/notebook_results")
code_box(page, "Model setup", "def build_dinov2_small():\n    return VisionTransformer(\n        image_size=224,\n        patch_size=14,\n        embed_dim=384,\n        depth=12,\n        num_heads=6,\n        layer_scale=1.0,\n    )", 7.0, 1.55, 5.65, 3.6, TEAL)
text(page, "Best checkpoint: 43/45 = 95.56%.", 7.08, 5.6, 5.35, .75, 15, NAVY, True)

# Slide 7
page = slide("DINOv2 fine-tuning", "DINOv2")
add_curve(page, "05_dinov2_fine_tuned/notebook_results_cpu_notebook")
code_box(page, "Load weights", "saved_state = torch.load(checkpoint, weights_only=True)\nsaved_state.pop(\"head.weight\", None)\nsaved_state.pop(\"head.bias\", None)\nmodel.load_state_dict(saved_state, strict=False)\n\n# Remove old ImageNet head.\n# Add Linear(384, 2) for Low / High.", 7.0, 1.55, 5.65, 3.75, ORANGE)
text(page, "Best checkpoint: 43/45 = 95.56%.", 7.08, 5.65, 5.4, .7, 15, NAVY, True)

# Slide 8
page = slide("VMamba frozen", "Visual Mamba")
add_curve(page, "06_vmamba_frozen/notebook_results")
code_box(page, "Classifier wrapper", "class VMambaClassifier(nn.Module):\n    def __init__(self, backbone):\n        self.backbone = backbone\n        self.head = nn.Linear(768, 2)\n\n    def forward(self, images):\n        features = self.backbone(images)\n        return self.head(features)", 7.0, 1.55, 5.65, 3.45, TEAL)
text(page, "Best checkpoint: 42/45 = 93.33%.",
     7.08, 5.55, 5.35, .8, 14, NAVY, True)

# Slide 9
page = slide("VMamba fine-tuning", "Visual Mamba")
add_curve(page, "07_vmamba_fine_tuned/notebook_results_fixed_cpu")
code_box(page, "Fine-tuning part", "model, last_stage = build_vmamba()\n\nconfigure_parameters(\n    model,\n    phase=\"fine_tuned\",\n    last_stage=last_stage,\n)\n\n# Only last_stage and Linear(768, 2) update.", 7.0, 1.55, 5.65, 3.65, ORANGE)
text(page, "Best checkpoint: 44/45 = 97.78%.",
     7.08, 5.65, 5.35, .7, 15, NAVY, True)

# Slide 10
page = slide("Checkpoint meaning", "checkpoints")
card(page, .75, 1.45, 3.72, 3.45, "During training", "Validation accuracy is checked after every epoch.", BLUE)
card(page, 4.81, 1.45, 3.72, 3.45, "Best checkpoint", "The best weights are saved and loaded again for the reported result.", TEAL)
card(page, 8.87, 1.45, 3.72, 3.45, "Limit", "This is still validation accuracy, not a separate test set.", ORANGE)
text(page, "For a stronger claim, we need a separate test set or repeated cross-validation.",
     .9, 5.95, 11.45, .45, 16, MUTED, True, align=PP_ALIGN.CENTER)

output = OUT / "notebook_runs_and_local_model_code_20260831.pptx"
prs.save(output)
print(output)
