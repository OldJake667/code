"""Create comparison graphs and a concise presentation from completed runs."""

import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
PRESENTATION = ROOT / "presentation"
PRESENTATION.mkdir(exist_ok=True)

EXPERIMENTS = [
    ("CNN", "01_cnn_from_scratch"),
    ("ViT-Tiny\nfrozen", "02_vit_tiny_frozen"),
    ("ViT-Tiny\nfine-tuned", "03_vit_tiny_fine_tuned"),
    ("DINOv2\nfrozen", "04_dinov2_frozen"),
    ("DINOv2\nfine-tuned", "05_dinov2_fine_tuned"),
    ("VMamba\nfrozen", "06_vmamba_frozen"),
    ("VMamba\nfine-tuned", "07_vmamba_fine_tuned"),
]


def read_results():
    rows = []
    for display, folder in EXPERIMENTS:
        result_file = ROOT / folder / "results/result_summary.json"
        history_file = ROOT / folder / "results/training_history.csv"
        if not result_file.is_file() or not history_file.is_file():
            raise FileNotFoundError(f"Experiment has not finished: {folder}")
        result = json.loads(result_file.read_text())
        with history_file.open(newline="", encoding="utf-8") as file:
            history = list(csv.DictReader(file))
        best_accuracy_value = max(float(row["validation_accuracy"]) for row in history)
        best_accuracy_epoch = next(
            int(row["epoch"])
            for row in history
            if float(row["validation_accuracy"]) == best_accuracy_value
        )
        best_loss_value = min(float(row["validation_loss"]) for row in history)
        best_loss_epoch = next(
            int(row["epoch"])
            for row in history
            if float(row["validation_loss"]) == best_loss_value
        )
        rows.append({
            "display": display,
            "folder": folder,
            "final_accuracy": result["final"]["validation_accuracy"],
            "final_correct": result["final"]["correct_predictions"],
            "best_accuracy": result["best_accuracy"]["validation_accuracy"],
            "best_accuracy_correct": result["best_accuracy"]["correct_predictions"],
            "best_accuracy_epoch": best_accuracy_epoch,
            "best_loss_accuracy": result["best_loss"]["validation_accuracy"],
            "best_loss_epoch": best_loss_epoch,
            "history": history,
        })
    return rows


rows = read_results()

with (ROOT / "comparison_results.csv").open("w", newline="", encoding="utf-8") as file:
    writer = csv.writer(file)
    writer.writerow([
        "experiment", "final_validation_accuracy", "final_correct_out_of_45",
        "best_validation_accuracy", "best_accuracy_correct_out_of_45",
        "best_accuracy_epoch", "lowest_loss_epoch",
    ])
    for row in rows:
        writer.writerow([
            row["display"].replace("\n", " "), row["final_accuracy"], row["final_correct"],
            row["best_accuracy"], row["best_accuracy_correct"],
            row["best_accuracy_epoch"], row["best_loss_epoch"],
        ])


def bar_chart(field, title, filename, color):
    labels = [row["display"] for row in rows]
    values = [100 * row[field] for row in rows]
    figure, axis = plt.subplots(figsize=(12, 5.4))
    bars = axis.bar(range(7), values, color=color)
    axis.set_xticks(range(7), labels)
    axis.set_ylim(75, 102)
    axis.set_ylabel("Accuracy (%)")
    axis.set_title(title, weight="bold")
    axis.grid(axis="y", alpha=0.22)
    for bar, value in zip(bars, values):
        axis.text(bar.get_x() + bar.get_width() / 2, value + 0.6,
                  f"{value:.2f}%", ha="center", fontsize=10, weight="bold")
    figure.tight_layout()
    figure.savefig(PRESENTATION / filename, dpi=220, facecolor="white")
    plt.close(figure)


bar_chart("final_accuracy", "Final scheduled-epoch validation accuracy",
          "final_validation_accuracy.png", "#25639B")
bar_chart("best_accuracy", "Highest observed validation accuracy",
          "best_validation_accuracy.png", "#1B9190")


figure, axes = plt.subplots(2, 4, figsize=(14, 7.2), sharey=True)
axes = axes.flatten()
for axis, row in zip(axes, rows):
    history = row["history"]
    epochs = [int(x["epoch"]) for x in history]
    train = [float(x["train_accuracy"]) for x in history]
    validation = [float(x["validation_accuracy"]) for x in history]
    axis.plot(epochs, train, color="#25639B", label="Training")
    axis.plot(epochs, validation, color="#1B9190", label="Validation")
    axis.set_title(row["display"].replace("\n", " "))
    axis.set_ylim(0.35, 1.03)
    axis.grid(alpha=0.2)
    axis.set_xlabel("Epoch")
    axis.set_ylabel("Accuracy")
axes[0].legend(frameon=False)
axes[-1].axis("off")
figure.tight_layout()
figure.savefig(PRESENTATION / "all_training_curves.png", dpi=200, facecolor="white")
plt.close(figure)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

BG, WHITE, NAVY = "F3F7FA", "FFFFFF", "0E1F32"
TEXT, MUTED, BLUE, TEAL, ORANGE = "263746", "607182", "25639B", "1B9190", "E17E2D"


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, text, x, y, w, h, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    run = paragraph.add_run(); run.text = text
    run.font.name = "Aptos"; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def new_slide(section, title):
    slide = prs.slides.add_slide(blank)
    background = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid(); background.fill.fore_color.rgb = rgb(BG); background.line.fill.background()
    slide.shapes._spTree.remove(background._element); slide.shapes._spTree.insert(2, background._element)
    add_text(slide, section.upper(), .68, .22, 12, .3, 11, TEAL, True)
    add_text(slide, title, .68, .56, 12, .72, 28, NAVY, True)
    return slide


def card(slide, x, y, w, h, heading, body, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb("D9E2E8")
    strip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(.08), Inches(h))
    strip.fill.solid(); strip.fill.fore_color.rgb = rgb(accent); strip.line.fill.background()
    add_text(slide, heading, x+.28, y+.24, w-.5, .4, 17, NAVY, True)
    add_text(slide, body, x+.28, y+.82, w-.5, h-1.0, 14, TEXT)


# Slide 1
slide = new_slide("MODEL REPRODUCTION", "Seven controlled image-classification experiments")
add_text(slide, "CNN  •  ViT-Tiny  •  DINOv2  •  VMamba", .75, 1.75, 11.8, .6, 24, BLUE, True)
add_text(slide, "Frozen-backbone and last-stage fine-tuning experiments use the same 176 training and 45 validation designs.",
         .75, 2.6, 11.5, 1.0, 20, TEXT)
add_text(slide, "Models are evaluated as classifiers of Low versus High bandwidth geometry images.",
         .75, 4.35, 11.5, .7, 18, MUTED)

# Slide 2
slide = new_slide("EXPERIMENT", "One controlled split—and no independent test set")
card(slide, .7, 1.55, 3.7, 3.6, "176 training images", "Used to update model parameters.", BLUE)
card(slide, 4.8, 1.55, 3.7, 3.6, "45 validation images", "Used during training and checkpoint selection. The filenames are identical for all seven runs.", TEAL)
card(slide, 8.9, 1.55, 3.7, 3.6, "0 independent test images", "A test accuracy cannot be reported without reserving a new untouched test set.", ORANGE)
add_text(slide, "One changed prediction changes accuracy by 2.22 percentage points.", .8, 5.75, 11.7, .5, 17, MUTED, True, PP_ALIGN.CENTER)

# Slide 3
slide = new_slide("IMPLEMENTATION", "What was constructed locally")
card(slide, .7, 1.5, 3.75, 4.6, "CNN", "Written layer by layer:\n\nConv2D → pooling → Conv2D → pooling → Conv2D → pooling → Dense(128) → Dense(2)", BLUE)
card(slide, 4.78, 1.5, 3.75, 4.6, "ViT-Tiny and DINOv2", "Patch embedding, class token, position embedding, attention, MLP, residual connections and LayerNorm are explicitly constructed in PyTorch.\n\nTheir loaded outputs exactly match timm.", TEAL)
card(slide, 8.86, 1.5, 3.75, 4.6, "VMamba", "Uses the model authors’ complete local implementation. A shortened rewrite would change selective scan and would not be the same pretrained architecture.", ORANGE)

# Slide 4
slide = new_slide("TRAINING", "Frozen and fine-tuned runs answer different questions")
card(slide, .9, 1.55, 5.5, 3.9, "Frozen backbone", "The pretrained feature extractor does not change. Only the new Low/High classifier learns.\n\n20 epochs · learning rate 0.0001", BLUE)
card(slide, 6.9, 1.55, 5.5, 3.9, "Last-stage fine-tuning", "Starts from the selected frozen checkpoint. The classifier and final block or stage are updated.\n\n10 additional epochs · learning rate 0.00001", TEAL)
add_text(slide, "Fine-tuning is a separate experiment; it does not replace the frozen result.", 1.0, 5.9, 11.3, .5, 17, MUTED, True, PP_ALIGN.CENTER)

# Slide 5
slide = new_slide("CHECKPOINTS", "Three model states are preserved")
card(slide, .75, 1.55, 3.75, 3.8, "Final epoch", "The model after the predeclared training schedule. This is the main final-epoch comparison.", BLUE)
card(slide, 4.8, 1.55, 3.75, 3.8, "Lowest validation loss", "Used to initialize the fine-tuning stage. Loss distinguishes epochs even when accuracy is tied.", TEAL)
card(slide, 8.85, 1.55, 3.75, 3.8, "Highest validation accuracy", "Saved separately as the best observed validation result. It is not an independent test result.", ORANGE)
add_text(slide, "After checkpoint selection, a new untouched test set would be needed to estimate final generalization.",
         .9, 5.85, 11.5, .7, 16, MUTED, True, PP_ALIGN.CENTER)

# Slides 6–8
slide = new_slide("RESULTS", "Final scheduled-epoch validation accuracy")
slide.shapes.add_picture(str(PRESENTATION / "final_validation_accuracy.png"), Inches(.65), Inches(1.35), width=Inches(12.0))

slide = new_slide("SUPPLEMENTARY RESULT", "Highest observed validation accuracy")
slide.shapes.add_picture(str(PRESENTATION / "best_validation_accuracy.png"), Inches(.65), Inches(1.35), width=Inches(12.0))
add_text(slide, "These values are selected on the validation set and should not be called test accuracy.", .8, 6.82, 11.7, .35, 12, MUTED, True, PP_ALIGN.CENTER)

slide = new_slide("LEARNING CURVES", "Training and validation accuracy across epochs")
slide.shapes.add_picture(str(PRESENTATION / "all_training_curves.png"), Inches(.65), Inches(1.30), width=Inches(12.0))

# Slide 9
slide = new_slide("INTERPRETATION", "What the results support—and what they do not")
final_correct_maximum = max(row["final_correct"] for row in rows)
final_ties = [row for row in rows if row["final_correct"] == final_correct_maximum]
final_maximum = final_correct_maximum / 45
final_names = " and ".join(row["display"].replace("\n", " ") for row in final_ties)
best_correct_maximum = max(row["best_accuracy_correct"] for row in rows)
best_ties = [row for row in rows if row["best_accuracy_correct"] == best_correct_maximum]
best_maximum = best_correct_maximum / 45
best_names = " and ".join(row["display"].replace("\n", " ") for row in best_ties)
card(slide, .75, 1.5, 3.75, 4.55, "Main comparison",
     f"Highest final-epoch result (tie):\n\n{final_names}\n{final_ties[0]['final_correct']}/45 = {100*final_maximum:.2f}%", BLUE)
card(slide, 4.8, 1.5, 3.75, 4.55, "Best observed",
     f"Highest selected validation result (tie):\n\n{best_names}\n{best_ties[0]['best_accuracy_correct']}/45 = {100*best_maximum:.2f}%", TEAL)
card(slide, 8.85, 1.5, 3.75, 4.55, "Careful conclusion",
     "The dataset is small. A one-image difference is 2.22 percentage points. These runs compare behavior on this split; they do not prove general architectural superiority or physical bandwidth improvement.", ORANGE)

output = PRESENTATION / "seven_model_code_and_results_20260831.pptx"
prs.save(output)
print(output)
